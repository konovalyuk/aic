import os
from flask import Flask, session, render_template, request, send_file, redirect, url_for, jsonify
from jinja2 import Template
from werkzeug.utils import secure_filename
from llama_models.llama3.api.datatypes import UserMessage, CompletionMessage, StopReason
from app.ai import load_generator, truncate_dialog
from app.extract_placeholders import extract_placeholders
from app.pdf import create_pdf_with_unicode
from config import ModelConfig, AppConfig
from PyPDF2 import PdfReader
import logging
import docx
from app.chat_completion import chat
from weasyprint import HTML

# Init Flask application
app = Flask(__name__)
app.config.from_object(AppConfig)

# Init logging
logging.basicConfig(level=logging.DEBUG)

generator = load_generator(ckpt_dir=ModelConfig.CKPT_DIR, model_parallel_size=ModelConfig.MODEL_PARALLEL_SIZE)


@app.route('/')
def home():
    """Initial page."""
    session.clear()
    return render_template('home.html')


@app.route('/upload', methods=['POST', 'GET'])
def upload():
    """Upload template of contract (PDF или DOCX)."""
    if request.method == 'POST':
        file = request.files['pdf']  # Имя должно совпадать с name="pdf" в форме
        if file and (file.filename.endswith('.pdf') or file.filename.endswith('.docx')):
            filename = secure_filename(file.filename)
            if '..' in filename or '/' in filename:
                return "Invalid file name."
            filepath = os.path.join(AppConfig.UPLOAD_FOLDER, filename)
            file.save(filepath)
            logging.info(f"Uploaded file saved to: {filepath}")
            return redirect(url_for('generate', filename=filename))
        return "Invalid file format. Please upload a PDF or DOCX."
    return render_template('upload.html')


@app.route('/generate/<filename>', methods=['GET', 'POST'])
def generate(filename):
    """Generate contract based on template."""
    if request.method == 'POST':
        description = request.form['description']
        filepath = os.path.join(AppConfig.UPLOAD_FOLDER, filename)

        # Чтение шаблона
        if filename.endswith('.pdf'):
            reader = PdfReader(filepath)
            text = " ".join([page.extract_text() for page in reader.pages])
        elif filename.endswith('.docx'):
            doc = docx.Document(filepath)
            text = " ".join([paragraph.text for paragraph in doc.paragraphs])

        # Генерация контракта
        dialog = [
            UserMessage(role="user", content=f"Here's a contract template: {text}"),
            UserMessage(role="user", content=f"Create a contract based on this description: {description}")
        ]
        dialog = truncate_dialog(dialog, ModelConfig.MAX_DIALOG_LENGTH)

        result = generator.chat_completion(dialog, max_gen_len=256)
        generated_content = result.generation.content

        output_path = os.path.join(AppConfig.UPLOAD_FOLDER, 'generated_contract.pdf')

        try:
            logging.info(f"Creating PDF at {output_path}")
            create_pdf_with_unicode(generated_content, output_path)
            logging.info(f"PDF created successfully at {output_path}")
        except Exception as e:
            logging.error("Error while creating PDF", exc_info=e)
            return "Failed to create PDF."

        if os.path.exists(output_path):
            return send_file(output_path, as_attachment=True)
        else:
            return "Generated file not found."

    return render_template('generate.html', filename=filename)


@app.route('/dialog', methods=['GET', 'POST'])
def dialog():
    """Диалоговое окно для взаимодействия с AI."""
    if request.method == 'POST':
        user_message = request.json.get('message', '')
        session_id = request.json.get('session_id', 'default_session')
        selected_model = request.json.get('model', '')

        if not user_message:
            return jsonify({'error': 'Message cannot be empty.'}), 400

        if not selected_model:
            return jsonify({'error': 'AI model not selected.'}), 400

        print("Raw session data:", session)

        # Получение или инициализация диалога в сессии
        dialog = session.get(session_id, [])
        dialog = [
            (UserMessage(**msg) if msg['role'] == 'user' else CompletionMessage(
                **{**msg, "stop_reason": StopReason(msg["stop_reason"])}  # Преобразование строки в StopReason
            ))
            for msg in dialog
        ]
        print("Deserialized dialog:", dialog)

        try:
            # Get assistant message
            assistant_message = chat(generator, dialog, user_message)
            dialog.append(assistant_message)

            session[session_id] = [
                {
                    **msg.dict(),
                    "stop_reason": msg.stop_reason.value  # Преобразование StopReason в строку
                } if isinstance(msg, CompletionMessage) else msg.dict()
                for msg in dialog
            ]

            return jsonify({'message': assistant_message.content})

        except Exception as e:
            logging.error("Error during dialog processing", exc_info=e)
            return jsonify({'error': 'Failed to generate a response.'}), 500

    return render_template('dialog.html')


@app.route('/generate_pdf', methods=['POST'])
def generate_pdf():
    """Завершает диалог и формирует PDF."""
    session_id = "unique-session-id"  # Используем фиксированный ID для примера
    dialog = session.get(session_id, [])

    if not dialog:
        return jsonify({'error': 'No dialog data available.'}), 400

    # Находим последнее сообщение от AI
    last_ai_message = next(
        (message['content'] for message in reversed(dialog) if message['role'] == 'assistant'),
        None
    )

    if not last_ai_message:
        return jsonify({'error': 'No AI message available to save.'}), 400

    print("!!! PDF last_ai_message:", last_ai_message)
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Добавляем только контент последнего сообщения от AI
    pdf.multi_cell(0, 10, last_ai_message)
    pdf.ln()

    pdf_path = os.path.join(AppConfig.UPLOAD_FOLDER, "contract.pdf")
    try:
        pdf.output(pdf_path)
        return jsonify({'pdf_url': f"/download/{os.path.basename(pdf_path)}"})
    except Exception as e:
        logging.error("Error while generating PDF", exc_info=e)
        return jsonify({'error': 'Failed to create PDF.'}), 500


@app.route('/download/<filename>')
def download_file(filename):
    """
    Позволяет скачать PDF-файл.
    """
    return send_file(os.path.join(AppConfig.UPLOAD_FOLDER, filename), as_attachment=True)


# ------------------------------------

@app.route('/custom_upload', methods=['GET', 'POST'])
def custom_upload():
    """Обработка загрузки HTML-шаблона контракта."""
    if request.method == 'POST':
        uploaded_file = request.files['html']
        if uploaded_file and uploaded_file.filename.endswith('.html'):
            file_path = os.path.join(AppConfig.UPLOAD_FOLDER, uploaded_file.filename)
            uploaded_file.save(file_path)

            # Используем Llama-модель для извлечения placeholders
            placeholders = extract_placeholders(file_path)

            # Сохраняем placeholders в сессии
            session['placeholders'] = placeholders
            session['template_path'] = file_path

            return redirect(url_for('custom_dialog'))

    return render_template('custom_upload.html')


@app.route('/custom_dialog', methods=['GET', 'POST'])
def custom_dialog():
    if request.method == 'GET':
        placeholders = session.get('placeholders', [])
        return render_template('custom_dialog.html', placeholders=placeholders)

    print("Raw session data:", session)
    data = request.get_json()
    user_message = data.get('message')
    session_id = data.get('session_id')
    model = data.get('model')
    print("!!! custom_dialog. data:", data)

    # Генерация ответа AI
    response_message = generate_ai_response(user_message, model, session_id)
    # response_message = ""

    return response_message
    # return jsonify({'message': response_message})


def generate_ai_response(user_message, model, session_id):
    if not user_message:
        return jsonify({'error': 'Message cannot be empty.'}), 400

    if not model:
        return jsonify({'error': 'AI model not selected.'}), 400

    # Получение или инициализация диалога в сессии
    dialog = session.get(session_id, [])
    dialog = [
        (UserMessage(**msg) if msg['role'] == 'user' else CompletionMessage(
            **{**msg, "stop_reason": StopReason(msg["stop_reason"])}  # Преобразование строки в StopReason
        ))
        for msg in dialog
    ]
    print("Deserialized dialog:", dialog)

    try:
        # Get assistant message
        assistant_message = chat(generator, dialog, user_message)
        print("!!! assistant_message:", assistant_message)
        dialog.append(assistant_message)

        # Определение и сохранение значений placeholders
        placeholders = session.get('placeholders', [])
        for placeholder in placeholders:
            if placeholder in user_message:
                session[placeholder] = extract_placeholder_value(user_message, placeholder)

        session[session_id] = [
            {
                **msg.dict(),
                "stop_reason": msg.stop_reason.value  # Преобразование StopReason в строку
            } if isinstance(msg, CompletionMessage) else msg.dict()
            for msg in dialog
        ]

        return jsonify({'message': assistant_message.content})

    except Exception as e:
        logging.error("Error during dialog processing", exc_info=e)
        return jsonify({'error': 'Failed to generate a response.'}), 500


def extract_placeholder_value(user_message, placeholder):
    # Простейший пример выделения значений для placeholders
    # Можно улучшить, используя регулярные выражения или NLP-техники
    parts = user_message.split(placeholder)
    if len(parts) > 1:
        value = parts[1].strip().split()[0]  # Предполагаем, что значение идёт сразу после placeholder
        return value
    return ""


@app.route('/custom_generate_pdf', methods=['POST'])
def custom_generate_pdf():
    """Генерация PDF из заполненного шаблона."""
    template_path = session.get('template_path')
    placeholders = session.get('placeholders', [])
    data = request.get_json()  # Получаем данные из JSON-запроса

    print("!!! placeholders:", placeholders)
    print("!!! data:", data)

    # Проверка данных на соответствие placeholders
    values = {key: data.get(key, '') for key in placeholders}
    print("!!! values:", values)

    # Рендеринг шаблона
    with open(template_path, 'r', encoding='utf-8') as file:
        template_content = file.read()
    template = Template(template_content)
    rendered_html = template.render(values)

    # Генерация PDF из заполненного шаблона
    pdf_output_path = os.path.join(AppConfig.UPLOAD_FOLDER, "generated_contract.pdf")
    HTML(string=rendered_html).write_pdf(pdf_output_path)

    try:
        return jsonify({'pdf_url': f"/download/{os.path.basename(pdf_output_path)}"})
    except Exception as e:
        logging.error("Error while generating PDF", exc_info=e)
        return jsonify({'error': 'Failed to create PDF.'}), 500
